#!/usr/bin/env python3
"""
Step 10 - Rebuild the collection instrument from the documents you ACTUALLY hold.

Reads out/extract_log.csv, keeps only those firm-years in the workbook, re-keys the
RDQI Scoring sheet to the new row positions, and adds a Collection Status sheet
showing each firm's realised panel and any documents still needing attention.

Your original workbook is never modified: output goes to a new file.

Usage:
    python 10_build_collection_tool.py --xlsx Paper1_Collection_Instrument_v3_REVISED.xlsx
    python 10_build_collection_tool.py --xlsx <file> --min-years 5
    python 10_build_collection_tool.py --xlsx <file> --dry-run
"""
import argparse, shutil
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="C6E0B4")

SHORT_CHARS = 5000        # below this, likely not a real reserve disclosure


def norm(v):
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out", default="out")
    ap.add_argument("--min-years", type=int, default=5,
                    help="firms with fewer usable years are flagged for exclusion")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    src = Path(a.xlsx)
    log = pd.read_csv(Path(a.out) / "extract_log.csv")
    log["firm_id"] = log["firm_id"].map(norm)
    log["year"] = log["year"].map(norm)
    log = log.dropna(subset=["firm_id", "year"])

    ok = log[(log["status"] == "ok") & (log["section_found"].astype(str) == "True")]
    ok = ok[~ok["needs_ocr"].astype(str).str.lower().eq("true")]
    ok = ok[ok["section_chars"] >= 200]

    have = set(zip(ok["firm_id"], ok["year"]))
    per_firm = ok.groupby("firm_id").size()
    short_firms = set(per_firm[per_firm < a.min_years].index)

    print(f"Documents in log        : {len(log)}")
    print(f"Usable (section found)  : {len(ok)}  across {per_firm.size} firms")
    print(f"Firms under {a.min_years} years    : {len(short_firms)}")

    flagged = log[(log["status"] != "ok")
                  | (log["section_found"].astype(str) != "True")
                  | (log["needs_ocr"].astype(str).str.lower() == "true")
                  | (log["section_chars"] < SHORT_CHARS)]
    print(f"Documents needing review: {len(flagged)}")

    if a.dry_run:
        print("\nDRY RUN - nothing written.")
        print("\nFirms that would be dropped for too few years:")
        for f in sorted(short_firms):
            print(f"   {f}: {per_firm[f]} years")
        return

    dst = src.with_name(src.stem.replace("_v3_REVISED", "") + "_v4_COLLECTED.xlsx")
    shutil.copyfile(src, dst)
    wb = openpyxl.load_workbook(dst)
    de, rs = wb["Data Entry"], wb["RDQI Scoring"]

    keep, drop = [], []
    for r in range(4, de.max_row + 1):
        eid, yr = norm(de.cell(r, 1).value), norm(de.cell(r, 12).value)
        if eid is None or yr is None:
            continue
        if (eid, yr) in have and eid not in short_firms:
            keep.append(r)
        else:
            drop.append(r)

    # delete in contiguous blocks, bottom-up: one call per block instead of per
    # row, which matters on a 97-column sheet with over a thousand rows to remove
    blocks, start, prev = [], None, None
    for r in sorted(drop):
        if start is None:
            start = prev = r
        elif r == prev + 1:
            prev = r
        else:
            blocks.append((start, prev - start + 1))
            start = prev = r
    if start is not None:
        blocks.append((start, prev - start + 1))
    for first, count in sorted(blocks, reverse=True):
        de.delete_rows(first, count)
    print(f"  deleted {len(drop)} rows in {len(blocks)} blocks")

    for i in range(len(keep)):
        rs.cell(5 + i, 1).value = f"='Data Entry'!A{4+i}"
        rs.cell(5 + i, 2).value = f"='Data Entry'!L{4+i}"
    for row in range(5 + len(keep), rs.max_row + 1):
        for c in range(1, 35):
            rs.cell(row, c).value = None

    # ---------------- Collection Status sheet ----------------
    if "Collection Status" in wb.sheetnames:
        del wb["Collection Status"]
    cs = wb.create_sheet("Collection Status", 2)
    cs["A1"] = "COLLECTION STATUS — built from extract_log.csv"
    cs["A1"].font = Font(bold=True, size=14, color=NAVY)
    cs["A2"] = (f"{len(ok)} usable documents across {per_firm.size} firms. Firms with fewer than "
                f"{a.min_years} usable years are excluded per the stated minimum-years rule. "
                "Panels are UNBALANCED by design: firms enter at listing or first reserve "
                "declaration and exit at delisting or cessation. Re-run this script whenever "
                "you add documents.")
    cs["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    cs.merge_cells("A2:H2"); cs.row_dimensions[2].height = 52

    hdr = ["entity_id", "company", "regime", "first_year", "last_year",
           "years_held", "flagged_docs", "status"]
    for c, h in enumerate(hdr, 1):
        cell = cs.cell(4, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for c, w in enumerate([11, 34, 16, 11, 11, 11, 13, 26], 1):
        cs.column_dimensions[get_column_letter(c)].width = w

    names, regimes = {}, {}
    wb_src = openpyxl.load_workbook(src, data_only=True)["Data Entry"]
    for r in range(4, wb_src.max_row + 1):
        e = norm(wb_src.cell(r, 1).value)
        if e is not None:
            names.setdefault(e, wb_src.cell(r, 2).value)
            regimes.setdefault(e, wb_src.cell(r, 5).value)

    flag_counts = flagged.groupby("firm_id").size().to_dict()
    rr = 5
    for fid, n in per_firm.sort_values(ascending=False).items():
        yrs = sorted(ok[ok["firm_id"] == fid]["year"])
        excluded = fid in short_firms
        st = (f"EXCLUDE - under {a.min_years} years" if excluded
              else "include" if n >= 10 else "include - short panel")
        vals = [fid, names.get(fid), regimes.get(fid), yrs[0], yrs[-1], n,
                flag_counts.get(fid, 0), st]
        for c, v in enumerate(vals, 1):
            cell = cs.cell(rr, c, v)
            if excluded:
                cell.fill = RED
            elif n < 10:
                cell.fill = AMBER
            elif c == 8:
                cell.fill = GREEN
        rr += 1
    cs.freeze_panes = "A5"

    # ---------------- documents needing review ----------------
    if "Docs To Review" in wb.sheetnames:
        del wb["Docs To Review"]
    dr = wb.create_sheet("Docs To Review", 3)
    dr["A1"] = "DOCUMENTS NEEDING MANUAL REVIEW"
    dr["A1"].font = Font(bold=True, size=13, color=NAVY)
    dr["A2"] = ("A section under 5,000 characters found by cue_density usually means the firm "
                "reports resources only - run the reserve-declaration check before fixing the "
                "extraction. needs_ocr means a scanned PDF with no text layer.")
    dr["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    dr.merge_cells("A2:G2"); dr.row_dimensions[2].height = 44
    cols = ["firm_id", "year", "file", "section_found", "section_method",
            "needs_ocr", "section_chars"]
    for c, h in enumerate(cols, 1):
        cell = dr.cell(4, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        dr.column_dimensions[get_column_letter(c)].width = [11, 8, 26, 14, 20, 11, 14][c - 1]
    for i, (_, row) in enumerate(flagged.iterrows()):
        for c, h in enumerate(cols, 1):
            dr.cell(5 + i, c, row.get(h))
    dr.freeze_panes = "A5"

    wb.save(dst)
    print(f"\nData Entry rows kept    : {len(keep)}  (removed {len(drop)})")
    print(f"Saved -> {dst}")
    print("Open it in Excel once and save, so the RDQI Scoring keys cache their values")
    print("before you run 07_to_workbook.py.")


if __name__ == "__main__":
    main()
