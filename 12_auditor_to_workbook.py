#!/usr/bin/env python3
"""
Step 12 - Write the auditor results into the workbook's Data Entry sheet.

07_to_workbook.py fills the RDQI Scoring sheet. This fills the big4 column (and,
if present, an auditor-name column) on Data Entry, matching on entity_id and year.

Columns are located by HEADER NAME in row 3, not by position, so it still works if
the workbook layout changes.

By default it will NOT overwrite a cell that already holds a value. Note that in
v3/v4 the big4 column is prefilled with 0 throughout, which is indistinguishable
from a genuine "non-Big 4" code - use --treat-zero-as-blank so those prefilled
zeros are treated as empty and get written.

Usage:
    python 12_auditor_to_workbook.py --xlsx <workbook>.xlsx --out out --dry-run
    python 12_auditor_to_workbook.py --xlsx <workbook>.xlsx --out out --treat-zero-as-blank
    python 12_auditor_to_workbook.py --xlsx <workbook>.xlsx --out out --min-confidence high
"""
import argparse, shutil
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill

HEADER_ROW = 3
FILL = PatternFill("solid", fgColor="FFF2CC")


def norm(v):
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def find_col(ws, *names):
    for c in range(1, ws.max_column + 1):
        h = ws.cell(HEADER_ROW, c).value
        if isinstance(h, str) and h.strip().lower() in names:
            return c
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out", default="out")
    ap.add_argument("--sheet", default="Data Entry")
    ap.add_argument("--min-confidence", default="any", choices=["any", "high"])
    ap.add_argument("--treat-zero-as-blank", action="store_true")
    ap.add_argument("--overwrite-existing", action="store_true")
    ap.add_argument("--mark", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    src = Path(a.xlsx)
    aud = pd.read_csv(Path(a.out) / "auditor.csv")
    aud = aud[aud["auditor"].notna() & (aud["auditor"].astype(str) != "")]
    if a.min_confidence == "high":
        aud = aud[aud["confidence"] == "high"]
    if aud.empty:
        raise SystemExit("No usable rows in auditor.csv")
    aud["key"] = list(zip(aud["firm_id"].map(norm), aud["year"].map(norm)))
    lookup = {k: (r["auditor"], int(r["big4"])) for k, r in
              zip(aud["key"], aud.to_dict("records"))}

    wb = openpyxl.load_workbook(src)
    ws = wb[a.sheet]
    c_id = find_col(ws, "entity_id")
    c_yr = find_col(ws, "year", "fiscal_year")
    c_b4 = find_col(ws, "big4", "audq", "big_4")
    c_nm = find_col(ws, "auditor", "auditor_name", "audit_firm")
    if not all([c_id, c_yr, c_b4]):
        raise SystemExit(f"Could not locate columns (entity_id={c_id}, year={c_yr}, "
                         f"big4={c_b4}). Check header row {HEADER_ROW}.")
    print(f"entity_id=col {c_id}, year=col {c_yr}, big4=col {c_b4}, "
          f"auditor={'col ' + str(c_nm) if c_nm else 'not present'}")

    written = protected = unmatched = 0
    b4_count = 0
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        eid, yr = norm(ws.cell(r, c_id).value), norm(ws.cell(r, c_yr).value)
        if eid is None or yr is None:
            continue
        hit = lookup.get((eid, yr))
        if hit is None:
            unmatched += 1
            continue
        name, b4 = hit
        cur = ws.cell(r, c_b4).value
        blank = cur in (None, "") or (a.treat_zero_as_blank and cur in (0, "0"))
        if not blank and not a.overwrite_existing:
            protected += 1
            continue
        if not a.dry_run:
            ws.cell(r, c_b4).value = b4
            if c_nm:
                ws.cell(r, c_nm).value = name
            if a.mark:
                ws.cell(r, c_b4).fill = FILL
        written += 1
        b4_count += b4

    print(f"\nRows written            : {written}   (Big 4 = {b4_count}, "
          f"non-Big 4 = {written - b4_count})")
    print(f"Rows left (had a value) : {protected}")
    print(f"Sheet rows with no match: {unmatched}")

    if a.dry_run:
        print("\nDRY RUN - nothing written.")
        return

    dst = src.with_name(src.stem + "_aud.xlsx")
    shutil.copyfile(src, dst)
    wb.save(dst)
    print(f"\nSaved -> {dst}")
    print("Unmatched rows are firm-years with no auditor identified - usually a")
    print("standalone reserve statement, which carries no auditor's report. Fill")
    print("those from the annual report or from S&P before using big4 in a model.")


if __name__ == "__main__":
    main()
