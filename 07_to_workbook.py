#!/usr/bin/env python3
"""
Step 7 - Write pipeline scores into the Paper 1 Collection Instrument workbook.

Fills columns C:Z (I1..I24) on the 'RDQI Scoring' sheet, matching each row by the
entity_id and year already in columns A and B. It does NOT touch the formula
columns (rdqi_total, rdqi_max, rdqi, rdqi_d1..d5) - those recalculate themselves
when you open the file in Excel.

It never overwrites your original workbook: output goes to a new file.

Usage:
    python 07_to_workbook.py --xlsx Paper1_Collection_Instrument_v2-4_9_26.xlsx
    python 07_to_workbook.py --xlsx <file> --overwrite-existing   # replace hand codes too
    python 07_to_workbook.py --xlsx <file> --dry-run              # report only, write nothing

Options:
    --mark      shade the cells this script wrote, so machine codes are visible
    --run       which scoring run to export (default 1)
"""
import argparse, shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def norm_id(v):
    """Entity ids arrive as 4188197, '4188197' or 4188197.0 depending on the
    source. Normalise all of them to the same string before matching."""
    s = str(v).strip()
    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
    except (TypeError, ValueError):
        pass
    return s


SHEET = "RDQI Scoring"
FIRST_COL, LAST_COL = 3, 26        # C..Z  =  I1..I24
HEADER_ROW = 4
FILL = PatternFill("solid", fgColor="FFF2CC")   # pale amber = machine coded


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out", default="out")
    ap.add_argument("--run", type=int, default=1)
    ap.add_argument("--sheet", default=SHEET)
    ap.add_argument("--overwrite-existing", action="store_true",
                    help="overwrite cells that already contain a score (e.g. hand codes)")
    ap.add_argument("--mark", action="store_true", help="shade cells this script wrote")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    src = Path(args.xlsx)
    scores = pd.read_csv(Path(args.out) / "quality_raw.csv")
    scores = scores[scores["run"] == args.run]
    if scores.empty:
        raise SystemExit(f"No rows for run {args.run} in quality_raw.csv")

    # --- map (entity_id, year) -> sheet row, using the COMPUTED values in A and B
    vals = load_workbook(src, data_only=True)[args.sheet]
    index = {}
    for r in range(HEADER_ROW + 1, vals.max_row + 1):
        eid, yr = vals.cell(r, 1).value, vals.cell(r, 2).value
        if eid in (None, "") or yr in (None, ""):
            continue
        try:
            index[(norm_id(eid), int(float(yr)))] = r
        except (TypeError, ValueError):
            continue
    if not index:
        raise SystemExit(
            "Could not read entity_id/year from columns A and B.\n"
            "Those cells are formulas: open the workbook in Excel once, press\n"
            "Ctrl+S to cache their values, close it, then run this again.")

    wb = load_workbook(src)
    ws = wb[args.sheet]
    items = [ws.cell(HEADER_ROW, c).value for c in range(FIRST_COL, LAST_COL + 1)]

    written = skipped = unmatched = protected = 0
    missing = []
    for _, row in scores.iterrows():
        key = (norm_id(row["firm_id"]), int(float(row["year"])))
        r = index.get(key)
        if r is None:
            unmatched += 1
            missing.append(key)
            continue
        for off, item in enumerate(items):
            col = FIRST_COL + off
            src_val = row.get("q_" + str(item))
            if pd.isna(src_val):
                skipped += 1
                continue
            cur = ws.cell(r, col).value
            if cur not in (None, "") and not args.overwrite_existing:
                protected += 1
                continue
            if not args.dry_run:
                v = str(src_val).strip()
                ws.cell(r, col).value = "NA" if v.upper() == "NA" else int(float(v))
                if args.mark:
                    ws.cell(r, col).fill = FILL
            written += 1

    print(f"Sheet rows indexed      : {len(index):,}")
    print(f"Firm-years to write     : {scores.shape[0]:,}")
    print(f"Cells written           : {written:,}")
    print(f"Cells left (had a value): {protected:,}   (use --overwrite-existing to replace)")
    print(f"Cells with no score     : {skipped:,}")
    print(f"Firm-years not matched  : {unmatched:,}")
    for k in missing[:10]:
        print(f"   no sheet row for entity {k[0]} year {k[1]}")
    if len(missing) > 10:
        print(f"   ... and {len(missing)-10} more")

    if args.dry_run:
        print("\nDRY RUN - nothing written.")
        return

    dst = src.with_name(src.stem + "_scored.xlsx")
    shutil.copyfile(src, dst)          # keep the original untouched
    wb.save(dst)
    print(f"\nSaved -> {dst}")
    print("Open it in Excel; rdqi_total / rdqi_max / rdqi / rdqi_d1..d5 recalculate on open.")
    print("IMPORTANT: always run this script against your ORIGINAL workbook, not")
    print("against a _scored file. Saving with openpyxl clears the cached values of")
    print("the entity_id and year formulas, so a _scored file cannot be matched again")
    print("until you have opened it in Excel and saved it there.")


if __name__ == "__main__":
    main()
